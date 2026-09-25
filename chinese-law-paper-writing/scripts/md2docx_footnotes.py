# -*- coding: utf-8 -*-
"""
md → docx（法学论文投稿稿）通用脚本
以样稿或期刊模板 docx 为模板继承样式与页面设置，写入新稿正文，
并用 zipfile 注入 Word 真实脚注（python-docx 原生不支持 footnotes part）。

两种交付模式：
  A. 干净版：不传 --red-* 参数，输出投稿用 docx；
  B. 修改标注版：传 --red-para-start / --red-footnote / --red-split-mark，
     新增或修改内容用红色字体，便于作者快速定位 AI 改动。

用法：
  python3 -X utf8 md2docx_footnotes.py --src template.docx --md draft.md --dst output.docx
  （依赖：python3 -m pip install python-docx lxml；Windows 用 py）

Markdown 约定：
- 标题可写 #/##/###/####，也可直接写“一、”“（一）”“1. ”；
- 正文脚注标记写 [N] 或 [脚注N]；文末 “## 脚注” 之后逐行写 “[N] 内容”；
- “摘　要：”“关键词：” 行按标签 + 内容双字体处理。

模板样式：优先使用模板中的“一级标题/二级标题/三级标题/正文1”，
缺失时回退到 Word 内置 Heading 1/2/3 与 Normal。字体字号集中在 FONTS 配置，
默认值只是中文法学期刊的常见版式，请按目标期刊要求或样稿实测修改。
插入的新内容须独立成段，否则按段首判定的红色标注会失效。
"""
import argparse
import re
import shutil
import tempfile
import zipfile
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

# ================= 配置区 =================
SRC = "/path/to/template.docx"   # 模板（继承样式/页边距）
MD  = "/path/to/draft.md"
DST = "/path/to/output.docx"
TMP = ""  # 留空时在输出目录旁自动创建并清理临时文件

# ---- 修改标注版配置（干净版留空即可） ----
RED_PARA_STARTS = ()
RED_FOOTNOTES = ()
RED_SPLIT_MARK = ""
RED_FOOTNOTE_ID = 0

# ---- 字体配置（东亚字体, 西文字体, 字号pt）；按目标期刊或样稿实测修改 ----
FONTS = {
    "title": ("黑体", "Times New Roman", 20),
    "author": ("宋体", "Times New Roman", 11),
    "affiliation": ("楷体", "楷体", 10.5),
    "abstract_label": ("黑体", "Times New Roman", 10.5),
    "abstract_body": ("楷体", "楷体", 10.5),
    "h1": ("黑体", "Times New Roman", 11),
    "h2": ("楷体", "楷体", 10.5),
    "h3": ("宋体", "宋体", 10.5),
    "body": ("宋体", "Times New Roman", 10.5),
    "footnote": ("宋体", "Times New Roman", 9),
}
STYLE_FALLBACKS = {
    "一级标题": "Heading 1",
    "二级标题": "Heading 2",
    "三级标题": "Heading 3",
    "正文1": "Normal",
}


def configure_cli():
    """Resolve paths from CLI while preserving the historical config block."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--src", help="template DOCX path; otherwise use SRC above")
    parser.add_argument("--md", dest="md_path", help="draft Markdown path; otherwise use MD above")
    parser.add_argument("--dst", dest="dst_path", help="output DOCX path; otherwise use DST above")
    parser.add_argument("--tmp", dest="tmp_path", help="optional working DOCX path")
    parser.add_argument("--red-para-start", action="append", default=None,
                        help="paragraph prefix to mark red; repeatable")
    parser.add_argument("--red-footnote", type=int, action="append", default=None,
                        help="footnote ID to mark red; repeatable")
    parser.add_argument("--red-split-mark", default=None, help="suffix marker to mark red in one footnote")
    parser.add_argument("--red-footnote-id", type=int, default=None)
    args = parser.parse_args()
    src = args.src or SRC
    md_path = args.md_path or MD
    dst_path = args.dst_path or DST
    if any(value.startswith("/path/to/") for value in (src, md_path, dst_path)):
        parser.error("provide --src, --md and --dst, or replace the SRC/MD/DST config values")
    src = os.path.abspath(os.path.expanduser(src))
    md_path = os.path.abspath(os.path.expanduser(md_path))
    dst_path = os.path.abspath(os.path.expanduser(dst_path))
    if not os.path.isfile(src):
        parser.error(f"template DOCX does not exist: {src}")
    if not os.path.isfile(md_path):
        parser.error(f"draft Markdown does not exist: {md_path}")
    if os.path.abspath(src) == dst_path:
        parser.error("output path must differ from the template DOCX")
    if not os.path.isdir(os.path.dirname(dst_path)):
        parser.error(f"output directory does not exist: {os.path.dirname(dst_path)}")
    if args.tmp_path:
        tmp_path = os.path.abspath(os.path.expanduser(args.tmp_path))
        remove_tmp = False
    else:
        fd, tmp_path = tempfile.mkstemp(prefix="md2docx-", suffix=".docx", dir=os.path.dirname(dst_path))
        os.close(fd)
        remove_tmp = True
    red_para_starts = tuple(args.red_para_start) if args.red_para_start is not None else RED_PARA_STARTS
    red_footnotes = tuple(args.red_footnote) if args.red_footnote is not None else RED_FOOTNOTES
    red_split_mark = args.red_split_mark if args.red_split_mark is not None else RED_SPLIT_MARK
    red_footnote_id = args.red_footnote_id if args.red_footnote_id is not None else RED_FOOTNOTE_ID
    return src, md_path, dst_path, tmp_path, remove_tmp, red_para_starts, red_footnotes, red_split_mark, red_footnote_id


SRC, MD, DST, TMP, REMOVE_TMP, RED_PARA_STARTS, RED_FOOTNOTES, RED_SPLIT_MARK, RED_FOOTNOTE_ID = configure_cli()

# ==========================================

shutil.copy(SRC, TMP)
doc = Document(TMP)
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

text = open(MD, encoding="utf-8").read()
body_md, _, fn_md = text.partition("## 脚注")
fn_map = {}
for m in re.finditer(r"^\[(\d+)\]\s*(.+)$", fn_md, re.M):
    fn_map[int(m.group(1))] = m.group(2).strip()

def set_run_font(run, east, west, size_pt, bold=False, red=False):
    run.font.name = west
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if red:
        run.font.color.rgb = RGBColor(0xFF, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), east)

_missing_styles = set()


def resolve_style(style_name):
    names = {style.name for style in doc.styles}
    if style_name in names:
        return doc.styles[style_name]
    fallback = STYLE_FALLBACKS.get(style_name)
    if style_name not in _missing_styles:
        _missing_styles.add(style_name)
        print(f"模板缺少样式“{style_name}”，回退为“{fallback or '默认段落'}”")
    if fallback and fallback in names:
        return doc.styles[fallback]
    return None


def add_para(style_name=None):
    p = doc.add_paragraph()
    if style_name:
        style = resolve_style(style_name)
        if style is not None:
            p.style = style
    return p


def add_run(p, text, font_key, red=False):
    east, west, size = FONTS[font_key]
    r = p.add_run(text)
    set_run_font(r, east, west, size, red=red)
    return r


def add_labeled(p, s, label_re):
    m = re.match(label_re, s)
    if m:
        add_run(p, m.group(1), "abstract_label")
        add_run(p, m.group(2), "abstract_body")
    else:
        add_run(p, s, "abstract_body")


first_content = True
for line in body_md.split("\n"):
    s = line.strip()
    if not s or s == "---":
        continue
    hashes = len(s) - len(s.lstrip("#"))
    if hashes and s[hashes:hashes + 1] == " ":
        s = s[hashes:].strip()
    else:
        hashes = 0
    if first_content and (hashes == 1 or (hashes == 0 and not s.startswith(("摘", "关键词")) and "作者姓名" not in s)):
        p = add_para(); p.alignment = 1
        add_run(p, s, "title")
        first_content = False
        continue
    first_content = False
    if s == "作者姓名":
        p = add_para(); p.alignment = 1
        add_run(p, s, "author")
    elif s.startswith("（单位名称"):
        p = add_para(); p.alignment = 1
        add_run(p, s, "affiliation")
    elif s.startswith(("摘　要", "摘要")):
        p = add_para(); p.alignment = 3
        add_labeled(p, s, r"^(摘\s*要[：:])(.*)$")
    elif s.startswith("关键词"):
        p = add_para(); p.alignment = 3
        add_labeled(p, s, r"^(关键词[：:])(.*)$")
    elif hashes == 2 or re.match(r"^[一二三四五六七八九十]+、", s):   # 一级标题
        p = add_para("一级标题")
        add_run(p, s, "h1")
    elif hashes == 3 or re.match(r"^（[一二三四五六七八九十]+）", s):  # 二级标题
        p = add_para("二级标题")
        add_run(p, s, "h2")
    elif hashes >= 4 or re.match(r"^\d+\.\s", s):                    # 三级标题
        p = add_para("三级标题")
        add_run(p, s, "h3")
    else:                                                               # 正文
        red_para = bool(RED_PARA_STARTS) and s.startswith(RED_PARA_STARTS)
        p = add_para("正文1")
        for idx, part in enumerate(re.split(r"\[(?:脚注)?(\d+)\]", s)):
            if idx % 2 == 0:
                if part:
                    add_run(p, part, "body", red=red_para)
            else:
                add_run(p, f"〔FNREF_{part}〕", "body", red=red_para)

doc.save(TMP)

# ---- 注入真实脚注（支持红色标注） ----
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
def w(tag):
    return f"{{{W}}}{tag}"

def make_text_run(content, red=False):
    r = etree.Element(w("r"))
    rPr = etree.SubElement(r, w("rPr"))
    fn_east, fn_west, fn_size = FONTS["footnote"]
    etree.SubElement(rPr, w("rFonts"), {w("ascii"): fn_west, w("hAnsi"): fn_west, w("eastAsia"): fn_east})
    half_points = str(int(round(fn_size * 2)))
    etree.SubElement(rPr, w("sz"), {w("val"): half_points})
    etree.SubElement(rPr, w("szCs"), {w("val"): half_points})
    if red:
        etree.SubElement(rPr, w("color"), {w("val"): "FF0000"})
    t = etree.SubElement(r, w("t"))
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = content
    return r

fn_root = etree.Element(w("footnotes"), nsmap={"w": W})
sep = etree.SubElement(fn_root, w("footnote"), {w("type"): "separator", w("id"): "-1"})
sp = etree.SubElement(sep, w("p")); spp = etree.SubElement(sp, w("pPr"))
etree.SubElement(spp, w("spacing"), {w("after"): "0", w("line"): "240", w("lineRule"): "auto"})
etree.SubElement(etree.SubElement(sp, w("r")), w("separator"))
cont = etree.SubElement(fn_root, w("footnote"), {w("type"): "continuationSeparator", w("id"): "0"})
cp = etree.SubElement(cont, w("p")); cpp = etree.SubElement(cp, w("pPr"))
etree.SubElement(cpp, w("spacing"), {w("after"): "0", w("line"): "240", w("lineRule"): "auto"})
etree.SubElement(etree.SubElement(cp, w("r")), w("continuationSeparator"))

for fid in sorted(fn_map.keys()):
    fn = etree.SubElement(fn_root, w("footnote"), {w("id"): str(fid)})
    p = etree.SubElement(fn, w("p"))
    etree.SubElement(etree.SubElement(p, w("pPr")), w("pStyle"), {w("val"): "FootnoteText"})
    r1 = etree.SubElement(p, w("r"))
    etree.SubElement(etree.SubElement(r1, w("rPr")), w("rStyle"), {w("val"): "FootnoteReference"})
    etree.SubElement(r1, w("footnoteRef"))
    content = fn_map[fid]
    if fid in RED_FOOTNOTES:
        p.append(make_text_run(content, red=True))
    elif fid == RED_FOOTNOTE_ID and RED_SPLIT_MARK and RED_SPLIT_MARK in content:
        head, _, tail = content.partition(RED_SPLIT_MARK)
        p.append(make_text_run(head))
        p.append(make_text_run(RED_SPLIT_MARK + tail, red=True))
    else:
        p.append(make_text_run(content))
footnotes_xml = etree.tostring(fn_root, xml_declaration=True, encoding="UTF-8", standalone=True)

zin = zipfile.ZipFile(TMP)
items = {}
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == "word/document.xml":
        xml = data.decode("utf-8")
        xml2, n = re.subn(r"〔FNREF_(\d+)〕",
                          lambda m: (f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
                                     f'<w:footnoteReference w:id="{m.group(1)}"/></w:r>'), xml)
        print(f"占位符替换: {n}")
        data = xml2.encode("utf-8")
    elif item.filename == "word/footnotes.xml":
        data = footnotes_xml
    items[item.filename] = (item, data)
zin.close()

# A minimal python-docx template may not already contain a footnotes part.
# Add the part plus its package relationship/content-type so the result remains
# a valid Word package instead of failing during post-generation verification.
if "word/footnotes.xml" not in items:
    items["word/footnotes.xml"] = (zipfile.ZipInfo("word/footnotes.xml"), footnotes_xml)

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

rels_name = "word/_rels/document.xml.rels"
if rels_name in items:
    rel_root = etree.fromstring(items[rels_name][1])
    if not any(node.get("Type") == REL_TYPE for node in rel_root):
        used_ids = {node.get("Id") for node in rel_root}
        rel_id = "rIdFootnotes"
        suffix = 2
        while rel_id in used_ids:
            rel_id = f"rIdFootnotes{suffix}"
            suffix += 1
        etree.SubElement(rel_root, f"{{{REL_NS}}}Relationship", {
            "Id": rel_id, "Type": REL_TYPE, "Target": "footnotes.xml"
        })
        items[rels_name] = (items[rels_name][0], etree.tostring(rel_root, xml_declaration=True, encoding="UTF-8"))

ct_name = "[Content_Types].xml"
if ct_name in items:
    ct_root = etree.fromstring(items[ct_name][1])
    footnote_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
    if not any(node.get("PartName") == "/word/footnotes.xml" for node in ct_root):
        etree.SubElement(ct_root, f"{{{CT_NS}}}Override", {
            "PartName": "/word/footnotes.xml", "ContentType": footnote_type
        })
        items[ct_name] = (items[ct_name][0], etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8"))

with zipfile.ZipFile(DST, "w", zipfile.ZIP_DEFLATED) as zout:
    for fname, (info, data) in items.items():
        zout.writestr(info, data)

with zipfile.ZipFile(DST) as z:
    docxml = z.read("word/document.xml").decode("utf-8")
    refs = re.findall(r'<w:footnoteReference w:id="(\d+)"', docxml)
    red_body = docxml.count('w:val="FF0000"')
    fnxml = z.read("word/footnotes.xml").decode("utf-8")
    red_fn = fnxml.count('w:val="FF0000"')
expected_ids = {str(fid) for fid in fn_map}
actual_ids = set(re.findall(r'<w:footnote w:id="(\d+)"', fnxml))
if set(refs) != expected_ids:
    raise RuntimeError(f"脚注引用与Markdown不一致: refs={sorted(set(refs))}, expected={sorted(expected_ids)}")
if not expected_ids.issubset(actual_ids):
    raise RuntimeError(f"footnotes.xml 缺少脚注条目: missing={sorted(expected_ids - actual_ids)}")
print(f"生成: {DST} | 脚注引用: {len(refs)} 唯一: {len(set(refs))} | "
      f"正文红色run: {red_body} 脚注红色run: {red_fn} | 大小: {os.path.getsize(DST)}")
if REMOVE_TMP:
    try:
        os.unlink(TMP)
    except FileNotFoundError:
        pass
