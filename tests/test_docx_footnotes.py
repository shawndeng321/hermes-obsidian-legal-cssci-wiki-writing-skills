"""End-to-end DOCX tests; every citation here is synthetic, not a source."""
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Mm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "law-paper-writing/scripts/md2docx_footnotes.py"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


class DocxFootnoteTests(unittest.TestCase):
    def setUp(self):
        self.work = tempfile.TemporaryDirectory(prefix="docx-footnote-test-")
        self.addCleanup(self.work.cleanup)
        self.root = Path(self.work.name)
        self.template = self.root / "template.docx"
        self.draft = self.root / "draft.md"
        self.output = self.root / "output.docx"
        Document().save(self.template)

    def convert(self, markdown, *options, template=True):
        self.draft.write_text(markdown, encoding="utf-8")
        args = [sys.executable, "-B", "-X", "utf8", str(SCRIPT)]
        if template:
            args.extend(["--src", str(self.template)])
        args.extend(["--md", str(self.draft), "--dst", str(self.output)])
        return subprocess.run(
            args + list(map(str, options)), cwd=ROOT, capture_output=True,
            text=True, encoding="utf-8", env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        )

    def xml(self, name):
        with zipfile.ZipFile(self.output) as archive:
            return etree.fromstring(archive.read(name))

    def assert_success(self, result):
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        Document(self.output)
        with zipfile.ZipFile(self.output) as archive:
            self.assertIsNone(archive.testzip())

    def refs(self):
        return self.xml("word/document.xml").xpath("//w:footnoteReference", namespaces=NS)

    def notes(self):
        return self.xml("word/footnotes.xml").xpath(
            "w:footnote[not(@w:type)]", namespaces=NS,
        )

    def text_runs(self, root):
        runs = []
        for run in root.xpath(".//w:r[w:t]", namespaces=NS):
            text = "".join(run.xpath("w:t/text()", namespaces=NS))
            italic = run.find("w:rPr/w:i", NS)
            color = run.find("w:rPr/w:color", NS)
            runs.append((text, italic is not None and italic.get(f"{{{W}}}val") != "0",
                         color.get(f"{{{W}}}val") if color is not None else None))
        return runs

    def test_standard_stable_id_becomes_a_real_word_footnote(self):
        result = self.convert(
            "# 合成占位标题\n\n占位正文[德]。[^S001]\n\n"
            "[^S001]: 合成引用占位，不是真实文献。\n"
        )
        self.assert_success(result)
        refs = self.refs()
        self.assertEqual(len(refs), 1)
        self.assertEqual(refs[0].getparent().tag, f"{{{W}}}r")
        self.assertEqual(refs[0].getparent().getparent().tag, f"{{{W}}}p")
        self.assertEqual(refs[0].get(f"{{{W}}}id"), "1")
        self.assertEqual(len(self.notes()), 1)
        self.assertEqual(self.notes()[0].get(f"{{{W}}}id"), "1")
        self.assertIn("合成引用占位，不是真实文献。", "".join(self.notes()[0].itertext()))
        body = "".join(self.xml("word/document.xml").itertext())
        self.assertIn("[德]", body)
        self.assertNotIn("[^S001]", body)
        self.assertNotIn("合成引用占位", body)

    def test_repeated_source_gets_one_word_note_per_occurrence(self):
        result = self.convert(
            "# 合成占位标题\n正文甲[^S001]；正文乙[^来源-甲]；再次引用[^S001]。\n\n"
            "[^来源-甲]: 合成来源乙占位。\n[^S001]: 合成来源甲占位。\n"
        )
        self.assert_success(result)
        ids = [ref.get(f"{{{W}}}id") for ref in self.refs()]
        self.assertEqual(ids, ["1", "2", "3"])
        notes = self.notes()
        self.assertEqual([note.get(f"{{{W}}}id") for note in notes], ids)
        self.assertEqual(
            ["".join(note.xpath(".//w:t/text()", namespaces=NS)) for note in notes],
            ["合成来源甲占位。", "合成来源乙占位。", "合成来源甲占位。"],
        )

    def test_explicit_emphasis_becomes_italic_without_guessing_english(self):
        result = self.convert(
            "# 合成占位标题\n正文*Synthetic Title*；Alpha _v._ Beta；Plain English[^S001]。\n"
            "[^S001]: [德]合成作者，*Synthetic Book*，Plain Publisher；Alpha _v._ Beta。\n"
        )
        self.assert_success(result)
        body_runs = self.text_runs(self.xml("word/document.xml"))
        note_runs = self.text_runs(self.notes()[0])
        self.assertIn(("Synthetic Title", True, None), body_runs)
        self.assertIn(("v.", True, None), body_runs)
        self.assertIn(("Synthetic Book", True, None), note_runs)
        self.assertIn(("v.", True, None), note_runs)
        self.assertTrue(any("Plain English" in text and not italic for text, italic, _ in body_runs))
        self.assertTrue(any("Plain Publisher" in text and not italic for text, italic, _ in note_runs))
        self.assertIn("[德]", "".join(self.notes()[0].itertext()))
        self.assertNotIn("*", "".join(self.notes()[0].itertext()))

    def test_red_source_selector_marks_all_repetitions_without_losing_italic(self):
        result = self.convert(
            "# 合成占位标题\n修订正文*Synthetic Term*[^S001]。\n再次引用[^S001]。\n"
            "[^S001]: 合成作者，*Synthetic Book*。\n",
            "--red-para-start", "修订正文", "--red-footnote", "S001",
        )
        self.assert_success(result)
        self.assertIn(("Synthetic Term", True, "FF0000"), self.text_runs(self.xml("word/document.xml")))
        self.assertEqual(len(self.notes()), 2)
        for note in self.notes():
            self.assertIn(("Synthetic Book", True, "FF0000"), self.text_runs(note))
            self.assertTrue(all(color == "FF0000" for _, _, color in self.text_runs(note)))
        color = self.refs()[0].getparent().find("w:rPr/w:color", NS)
        self.assertEqual(color.get(f"{{{W}}}val"), "FF0000")

    def test_red_suffix_can_split_inside_an_italic_span(self):
        result = self.convert(
            "# 合成占位标题\n正文[7]。\n## 脚注\n"
            "[7] 开头，*Synthetic Revised Title*，末尾。\n",
            "--red-footnote-id", "7", "--red-split-mark", "Revised",
        )
        self.assert_success(result)
        runs = self.text_runs(self.notes()[0])
        self.assertIn(("Synthetic ", True, None), runs)
        self.assertIn(("Revised Title", True, "FF0000"), runs)
        self.assertIn(("，末尾。", False, "FF0000"), runs)
        self.assertEqual("".join(text for text, _, _ in runs), "开头，Synthetic Revised Title，末尾。")

    def test_duplicate_definitions_are_rejected_before_output(self):
        for markdown in (
            "# 合成占位标题\n正文[^S001]。\n[^S001]: 合成甲。\n[^S001]: 合成乙。\n",
            "# 合成占位标题\n正文[1]。\n## 脚注\n[1] 合成甲。\n[1] 合成乙。\n",
        ):
            with self.subTest(markdown=markdown):
                result = self.convert(markdown)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("重复脚注定义", result.stderr)
                self.assertFalse(self.output.exists())

    def test_missing_definitions_are_rejected_before_output(self):
        for marker in ("[^S001]", "[9]", "[脚注9]"):
            with self.subTest(marker=marker):
                result = self.convert(f"# 合成占位标题\n正文{marker}。\n")
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("缺失脚注定义", result.stderr)
                self.assertFalse(self.output.exists())

    def test_unreferenced_definitions_are_rejected_before_output(self):
        for definition in ("[^orphan]: 合成孤立引用。", "## 脚注\n[3] 合成孤立引用。"):
            with self.subTest(definition=definition):
                result = self.convert(f"# 合成占位标题\n无引注占位正文。\n{definition}\n")
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("孤立脚注定义", result.stderr)
                self.assertFalse(self.output.exists())

    def test_existing_and_input_files_are_never_overwritten(self):
        markdown = "# 合成占位标题\n正文[^S001]。\n[^S001]: 合成引用。\n"
        self.output.write_bytes(b"existing-output-sentinel")
        result = self.convert(markdown)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("already exists", result.stderr)
        self.assertEqual(self.output.read_bytes(), b"existing-output-sentinel")
        self.output.unlink()
        working = self.root / "working.docx"
        working.write_bytes(b"existing-work-sentinel")
        for target in (working, self.template, self.draft, self.output):
            with self.subTest(target=target.name):
                self.draft.write_text(markdown, encoding="utf-8")
                before = target.read_bytes() if target.exists() else None
                result = self.convert(markdown, "--tmp", target)
                self.assertNotEqual(result.returncode, 0)
                if before is not None:
                    self.assertEqual(target.read_bytes(), before)
                self.assertFalse(self.output.exists())

    def test_no_template_requires_explicit_provisional_page_mode(self):
        result = self.convert(
            "# 合成占位标题\n正文[^S001]。\n[^S001]: 合成引用。\n",
            "--no-template", template=False,
        )
        self.assert_success(result)
        self.assertIn("临时页面", result.stdout)
        self.assertIn("暂定", result.stdout)
        self.assertIn("--src", result.stdout)
        section = Document(self.output).sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=0.1)
        self.assertAlmostEqual(section.page_height.mm, 297, delta=0.1)
        for value in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin):
            self.assertAlmostEqual(value.mm, 25.4, delta=0.1)

    def test_inline_footnotes_are_emitted_in_all_paragraph_roles(self):
        result = self.convert(
            "# 合成标题[^S001]\n摘要：合成摘要[^S001]\n关键词：合成词[^S001]\n"
            "## 合成一级标题[^S001]\n### 合成二级标题[^S001]\n#### 合成三级标题[^S001]\n"
            "[^S001]: 合成引用占位。\n"
        )
        self.assert_success(result)
        self.assertEqual([ref.get(f"{{{W}}}id") for ref in self.refs()], [str(i) for i in range(1, 7)])
        self.assertEqual(len(self.notes()), 6)
        self.assertNotIn("[^S001]", "".join(self.xml("word/document.xml").itertext()))

    def test_standard_definitions_work_under_a_legacy_note_heading(self):
        result = self.convert(
            "# 合成占位标题\n正文[^S001]与[脚注2]。\n\n## 脚注\n"
            "[^S001]: 合成标准引注。\n[2] 合成旧式引注。\n"
        )
        self.assert_success(result)
        self.assertEqual(len(self.refs()), 2)
        self.assertEqual(
            ["".join(note.xpath(".//w:t/text()", namespaces=NS)) for note in self.notes()],
            ["合成标准引注。", "合成旧式引注。"],
        )

    def test_indented_multiline_definition_stays_inside_the_word_note(self):
        result = self.convert(
            "# 合成占位标题\n正文[^S001]。\n\n"
            "[^S001]: 合成第一行\n    *Synthetic Continued Title*。\n\n"
            "    合成第二段。\n\n正常后续正文。\n"
        )
        self.assert_success(result)
        note = self.notes()[0]
        self.assertEqual(len(note.findall("w:p", NS)), 2)
        self.assertEqual(len(note.findall(".//w:footnoteRef", NS)), 1)
        self.assertIn(("Synthetic Continued Title", True, None), self.text_runs(note))
        self.assertIn("合成第一行 Synthetic Continued Title。", "".join(note.itertext()))
        body = "".join(self.xml("word/document.xml").itertext())
        self.assertNotIn("合成第二段", body)
        self.assertIn("正常后续正文", body)

    def test_note_heading_is_not_detected_inside_body_text(self):
        result = self.convert(
            "# 合成占位标题\n正文提到 ## 脚注 的占位文字[^S001]。\n"
            "[^S001]: 合成引用。\n"
        )
        self.assert_success(result)
        self.assertIn("正文提到 ## 脚注 的占位文字", "".join(self.xml("word/document.xml").itertext()))
        self.assertEqual(len(self.refs()), 1)

    def test_invalid_markdown_leaves_no_working_or_output_files(self):
        result = self.convert("# 合成标题\n正文[^missing]。\n")
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(sorted(path.name for path in self.root.iterdir()), ["draft.md", "template.docx"])
        working = self.root / "explicit-work.docx"
        result = self.convert("# 合成标题\n正文[^missing]。\n", "--tmp", working)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(working.exists())
        self.assertFalse(self.output.exists())

    def test_generated_footnote_style_references_resolve_to_real_styles(self):
        result = self.convert("# 合成标题\n正文[^S001]。\n[^S001]: 合成引用。\n")
        self.assert_success(result)
        styles = {
            style.get(f"{{{W}}}styleId"): style.get(f"{{{W}}}type")
            for style in self.xml("word/styles.xml").findall("w:style", NS)
        }
        paragraph = self.notes()[0].find("w:p/w:pPr/w:pStyle", NS)
        self.assertEqual(styles.get(paragraph.get(f"{{{W}}}val")), "paragraph")
        for root in (self.xml("word/document.xml"), self.notes()[0]):
            for reference_style in root.findall(".//w:rPr/w:rStyle", NS):
                self.assertEqual(styles.get(reference_style.get(f"{{{W}}}val")), "character")

    def test_template_section_and_named_paragraph_styles_are_preserved(self):
        template = Document()
        section = template.sections[0]
        section.page_width, section.page_height = Mm(180), Mm(260)
        section.left_margin, section.top_margin = Mm(31), Mm(23)
        section.header.paragraphs[0].text = "合成模板页眉"
        for name in ("一级标题", "二级标题", "三级标题", "正文1"):
            style = template.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.paragraph_format.space_after = Pt(7)
        template.add_paragraph("旧模板正文，必须移除")
        template.save(self.template)
        original_section = etree.tostring(template.element.body.sectPr)
        result = self.convert(
            "# 合成标题\n## 一级占位\n### 二级占位\n#### 三级占位\n正文[脚注8]。\n## 脚注\n[8] 合成引用。\n"
        )
        self.assert_success(result)
        generated = Document(self.output)
        self.assertEqual(etree.tostring(generated.element.body.sectPr), original_section)
        self.assertEqual([p.style.name for p in generated.paragraphs[1:]],
                         ["一级标题", "二级标题", "三级标题", "正文1"])
        self.assertEqual(generated.styles["正文1"].paragraph_format.space_after.pt, 7)
        self.assertEqual(generated.sections[0].header.paragraphs[0].text, "合成模板页眉")
        self.assertNotIn("旧模板正文", "".join(self.xml("word/document.xml").itertext()))
        run = generated.paragraphs[-1].runs[0]
        self.assertEqual(run.font.name, "Times New Roman")
        self.assertEqual(run.font.size.pt, 10.5)
        self.assertEqual(run._element.find("w:rPr/w:rFonts", NS).get(f"{{{W}}}eastAsia"), "宋体")

    def test_missing_heading_styles_fall_back_to_word_builtins(self):
        result = self.convert("# 合成标题\n## 一级\n### 二级\n#### 三级\n正文。\n")
        self.assert_success(result)
        self.assertEqual([p.style.name for p in Document(self.output).paragraphs[1:]],
                         ["Heading 1", "Heading 2", "Heading 3", "Normal"])
        self.assertIn("回退", result.stdout)

    def test_footnote_package_relationships_and_system_notes_are_complete(self):
        result = self.convert("# 合成标题\n正文[^S001]。\n[^S001]: 合成引用。\n")
        self.assert_success(result)
        rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
        rels = self.xml("word/_rels/document.xml.rels")
        footnote_rels = [rel for rel in rels if rel.get("Type") == rel_type]
        self.assertEqual(len(footnote_rels), 1)
        self.assertEqual(footnote_rels[0].get("Target"), "footnotes.xml")
        ct = self.xml("[Content_Types].xml")
        overrides = [entry for entry in ct if entry.get("PartName") == "/word/footnotes.xml"]
        self.assertEqual(len(overrides), 1)
        self.assertTrue(overrides[0].get("ContentType").endswith("footnotes+xml"))
        system = self.xml("word/footnotes.xml").xpath("w:footnote[@w:type]", namespaces=NS)
        self.assertEqual([(note.get(f"{{{W}}}id"), note.get(f"{{{W}}}type")) for note in system],
                         [("-1", "separator"), ("0", "continuationSeparator")])
        self.assertEqual(len(self.notes()[0].findall(".//w:footnoteRef", NS)), 1)

    def test_template_old_footnotes_are_replaced_without_orphans(self):
        result = self.convert("# 合成旧标题\n旧正文[^OLD]。\n[^OLD]: 合成旧引用必须移除。\n")
        self.assert_success(result)
        self.template.write_bytes(self.output.read_bytes())
        self.output.unlink()
        result = self.convert("# 合成新标题\n新正文[^NEW]。\n[^NEW]: 合成新引用。\n")
        self.assert_success(result)
        self.assertEqual(len(self.refs()), 1)
        self.assertEqual(len(self.notes()), 1)
        self.assertNotIn("合成旧引用", "".join(self.xml("word/footnotes.xml").itertext()))

    def test_stable_ids_are_opaque_to_the_emphasis_parser(self):
        result = self.convert(
            "# 合成标题\n正文[^_source_]与[^source*note*]；*Synthetic [^_source_] Span*。\n"
            "[^_source_]: 合成下划线引用。\n[^source*note*]: 合成星号引用。\n"
        )
        self.assert_success(result)
        self.assertEqual(len(self.refs()), 3)
        self.assertEqual(len(self.notes()), 3)
        body = self.text_runs(self.xml("word/document.xml"))
        self.assertIn(("Synthetic ", True, None), body)
        self.assertIn((" Span", True, None), body)
        self.assertNotIn("[^", "".join(self.xml("word/document.xml").itertext()))


if __name__ == "__main__":
    unittest.main()
