# -*- coding: utf-8 -*-
"""
md → docx（法学论文投稿稿）通用脚本
以样稿或期刊模板 docx 为模板继承样式与页面设置，写入新稿正文，
并用 zipfile 注入 Word 真实脚注（python-docx 原生不支持 footnotes part）。

两种交付模式：
  A. 干净版：不传 --red-* 参数，输出投稿用 docx；
  B. 修改标注版：传 --red-para-start / --red-footnote / --red-split-mark，
     新增或修改内容用红色字体，便于作者快速定位 AI 改动。

用法：
  python3 -X utf8 md2docx_footnotes.py --src template.docx --md draft.md --dst output.docx
  （依赖：python3 -m pip install python-docx lxml；Windows 用 py）

Markdown 约定：
- 标题可写 #/##/###/####，也可直接写“一、”“（一）”“1. ”；
- 正文脚注标记写 [N] 或 [脚注N]；文末 “## 脚注” 之后逐行写 “[N] 内容”；
- “摘　要：”“关键词：” 行按标签 + 内容双字体处理。

模板样式：优先使用模板中的“一级标题/二级标题/三级标题/正文1”，
缺失时回退到 Word 内置 Heading 1/2/3 与 Normal。字体字号集中在 FONTS 配置，
默认值只是中文法学期刊的常见版式，请按目标期刊要求或样稿实测修改。
插入的新内容须独立成段，否则按段首判定的红色标注会失效。
"""
import argparse
import re
import shutil
import tempfile
import zipfile
import os
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor, Mm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

# ================= 配置区 =================
SRC = "/path/to/template.docx"   # 模板（继承样式/页边距）
MD  = "/path/to/draft.md"
DST = "/path/to/output.docx"
TMP = ""  # 留空时在输出目录旁自动创建并清理临时文件

# ---- 修改标注版配置（干净版留空即可） ----
RED_PARA_STARTS = ()
RED_FOOTNOTES = ()
RED_SPLIT_MARK = ""
RED_FOOTNOTE_ID = 0

# ---- 字体配置（东亚字体, 西文字体, 字号pt）；按目标期刊或样稿实测修改 ----
FONTS = {
    "title": ("黑体", "Times New Roman", 20),
    "author": ("宋体", "Times New Roman", 11),
    "affiliation": ("楷体", "楷体", 10.5),
    "abstract_label": ("黑体", "Times New Roman", 10.5),
    "abstract_body": ("楷体", "楷体", 10.5),
    "h1": ("黑体", "Times New Roman", 11),
    "h2": ("楷体", "楷体", 10.5),
    "h3": ("宋体", "宋体", 10.5),
    "body": ("宋体", "Times New Roman", 10.5),
    "footnote": ("宋体", "Times New Roman", 9),
}
STYLE_FALLBACKS = {
    "一级标题": "Heading 1",
    "二级标题": "Heading 2",
    "三级标题": "Heading 3",
    "正文1": "Normal",
}


def configure_cli():
    """Resolve paths from CLI while preserving the historical config block."""
    parser = argparse.ArgumentParser(description=__doc__)
    template_mode = parser.add_mutually_exclusive_group()
    template_mode.add_argument("--src", help="template DOCX path; otherwise use SRC above")
    template_mode.add_argument("--no-template", action="store_true",
                               help="explicitly accept reversible provisional A4 / 25.4 mm pages")
    parser.add_argument("--md", dest="md_path", help="draft Markdown path; otherwise use MD above")
    parser.add_argument("--dst", dest="dst_path", help="output DOCX path; otherwise use DST above")
    parser.add_argument("--tmp", dest="tmp_path", help="optional working DOCX path")
    parser.add_argument("--red-para-start", action="append", default=None,
                        help="paragraph prefix to mark red; repeatable")
    parser.add_argument("--red-footnote", action="append", default=None,
                        help="Markdown source ID to mark red at every occurrence; repeatable")
    parser.add_argument("--red-split-mark", default=None, help="suffix marker to mark red in one footnote")
    parser.add_argument("--red-footnote-id", default=None, help="Markdown source ID for the red suffix")
    args = parser.parse_args()
    src = None if args.no_template else args.src or SRC
    md_path = args.md_path or MD
    dst_path = args.dst_path or DST
    if any(value and value.startswith("/path/to/") for value in (src, md_path, dst_path)):
        parser.error("provide --src (or --no-template), --md and --dst, or replace the config values")
    src = os.path.abspath(os.path.expanduser(src)) if src else None
    md_path = os.path.abspath(os.path.expanduser(md_path))
    dst_path = os.path.abspath(os.path.expanduser(dst_path))
    if src and not os.path.isfile(src):
        parser.error(f"template DOCX does not exist: {src}")
    if not os.path.isfile(md_path):
        parser.error(f"draft Markdown does not exist: {md_path}")
    if src == dst_path:
        parser.error("output path must differ from the template DOCX")
    if dst_path == md_path:
        parser.error("output path must differ from the draft Markdown")
    if os.path.lexists(dst_path):
        parser.error(f"output already exists; choose a new versioned path: {dst_path}")
    if not os.path.isdir(os.path.dirname(dst_path)):
        parser.error(f"output directory does not exist: {os.path.dirname(dst_path)}")
    if args.tmp_path:
        tmp_path = os.path.abspath(os.path.expanduser(args.tmp_path))
        if tmp_path in (src, md_path, dst_path):
            parser.error("working path must differ from template, Markdown and output")
        if os.path.lexists(tmp_path):
            parser.error(f"working file already exists: {tmp_path}")
        if not os.path.isdir(os.path.dirname(tmp_path)):
            parser.error(f"working directory does not exist: {os.path.dirname(tmp_path)}")
        remove_tmp = False
    else:
        tmp_path = None
        remove_tmp = True
    red_para_starts = tuple(args.red_para_start) if args.red_para_start is not None else RED_PARA_STARTS
    red_footnotes = tuple(args.red_footnote) if args.red_footnote is not None else RED_FOOTNOTES
    red_split_mark = args.red_split_mark if args.red_split_mark is not None else RED_SPLIT_MARK
    red_footnote_id = args.red_footnote_id if args.red_footnote_id is not None else RED_FOOTNOTE_ID
    red_footnotes = tuple(str(key) for key in red_footnotes)
    red_footnote_id = str(red_footnote_id)
    return src, md_path, dst_path, tmp_path, remove_tmp, red_para_starts, red_footnotes, red_split_mark, red_footnote_id


SRC, MD, DST, TMP, REMOVE_TMP, RED_PARA_STARTS, RED_FOOTNOTES, RED_SPLIT_MARK, RED_FOOTNOTE_ID = configure_cli()

# ==========================================

text = open(MD, encoding="utf-8").read()
parts = re.split(r"(?m)^##[ \t]+脚注[ \t]*\r?$", text, maxsplit=1)
body_md, fn_md = parts[0], parts[1] if len(parts) > 1 else ""
source_notes = {}


def register_note(key, content):
    if key in source_notes:
        raise SystemExit(f"重复脚注定义: {key}")
    source_notes[key] = content.strip()


def extract_definitions(part, legacy_section=False):
    lines = part.splitlines()
    body_lines = []
    index = 0
    while index < len(lines):
        definition = re.match(r"^\[\^([^\]\s]+)\]:[ \t]*(.*)$", lines[index])
        if not definition and legacy_section:
            definition = re.match(r"^\[(\d+)\][ \t]*(.*)$", lines[index])
        if not definition:
            body_lines.append(lines[index])
            index += 1
            continue
        key, first_line = definition.groups()
        paragraphs = [first_line.strip()]
        index += 1
        while index < len(lines):
            if lines[index].startswith(("    ", "\t")):
                continuation = lines[index].strip()
                paragraphs[-1] += (" " if paragraphs[-1] else "") + continuation
                index += 1
            elif (not lines[index].strip() and index + 1 < len(lines)
                  and lines[index + 1].startswith(("    ", "\t"))):
                paragraphs.append("")
                index += 1
            else:
                break
        register_note(key, "\n\n".join(paragraphs))
    return "\n".join(body_lines)


extract_definitions(fn_md, legacy_section=True)
body_md = extract_definitions(body_md)
fn_map = {}
fn_sources = {}
REFERENCE_RE = re.compile(r"\[\^([^\]\s]+)\]|\[(?:脚注)?(\d+)\]")
referenced_sources = {match.group(1) or match.group(2) for match in REFERENCE_RE.finditer(body_md)}
missing_sources = referenced_sources - source_notes.keys()
if missing_sources:
    raise SystemExit(f"缺失脚注定义: {', '.join(sorted(missing_sources))}")
orphan_sources = source_notes.keys() - referenced_sources
if orphan_sources:
    raise SystemExit(f"孤立脚注定义: {', '.join(sorted(orphan_sources))}")

if TMP is None:
    fd, TMP = tempfile.mkstemp(prefix="md2docx-", suffix=".docx", dir=os.path.dirname(DST))
    os.close(fd)
else:
    with open(TMP, "xb"):
        pass
if SRC:
    shutil.copy(SRC, TMP)
    doc = Document(TMP)
else:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(25.4)
    section.left_margin = section.right_margin = Mm(25.4)
    print("未提供样稿：临时页面 A4（210 × 297 mm），四边页边距 25.4 mm；均为暂定，"
          "不代表期刊要求。取得模板后用 --src 重新生成新版本即可撤销。")
body = doc.element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

EMPHASIS_RE = re.compile(r"(?<!\*)\*([^*\n]+)\*(?!\*)|(?<![\w_])_([^_\n]+)_(?![\w_])")


def inline_spans(text):
    """Parse explicit single-marker emphasis, never infer foreign-language ranges."""
    end = 0
    references = [match.span() for match in REFERENCE_RE.finditer(text)]
    for match in EMPHASIS_RE.finditer(text):
        if any(start <= match.start() < stop or start < match.end() <= stop
               for start, stop in references):
            continue
        if match.start() > end:
            yield text[end:match.start()], False, end
        group = 1 if match.group(1) is not None else 2
        yield match.group(group), True, match.start(group)
        end = match.end()
    if end < len(text) or not text:
        yield text[end:], False, end

def set_run_font(run, east, west, size_pt, bold=False, red=False):
    run.font.name = west
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if red:
        run.font.color.rgb = RGBColor(0xFF, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), east)


def ensure_footnote_style(style_id, name, style_type):
    for style in doc.styles:
        if (style.style_id == style_id or style.name == name) and style.type == style_type:
            return style
    names = {style.name for style in doc.styles}
    if name in names or any(style.style_id == style_id for style in doc.styles):
        name = f"Generated {name}"
        while name in names:
            name = f"Generated {name}"
    style = doc.styles.add_style(name, style_type)
    if style_type == WD_STYLE_TYPE.CHARACTER:
        style.font.superscript = True
    else:
        style.base_style = doc.styles["Normal"]
    return style


FOOTNOTE_TEXT_STYLE = ensure_footnote_style("FootnoteText", "Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
FOOTNOTE_REF_STYLE = ensure_footnote_style("FootnoteReference", "Footnote Reference", WD_STYLE_TYPE.CHARACTER)

_missing_styles = set()


def resolve_style(style_name):
    names = {style.name for style in doc.styles}
    if style_name in names:
        return doc.styles[style_name]
    fallback = STYLE_FALLBACKS.get(style_name)
    if style_name not in _missing_styles:
        _missing_styles.add(style_name)
        print(f"模板缺少样式“{style_name}”，回退为“{fallback or '默认段落'}”")
    if fallback and fallback in names:
        return doc.styles[fallback]
    return None


def add_para(style_name=None):
    p = doc.add_paragraph()
    if style_name:
        style = resolve_style(style_name)
        if style is not None:
            p.style = style
    return p


def add_plain_run(p, content, font_key, red=False, italic=False):
    east, west, size = FONTS[font_key]
    run = p.add_run(content)
    set_run_font(run, east, west, size, red=red)
    run.font.italic = italic
    return run


def add_reference(p, key, font_key, red=False):
    fid = len(fn_map) + 1
    fn_map[fid] = source_notes[key]
    fn_sources[fid] = key
    run = add_plain_run(p, "", font_key, red=red)
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), FOOTNOTE_REF_STYLE.style_id)
    run._element.get_or_add_rPr().insert(0, rStyle)
    run.font.superscript = True
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), str(fid))
    run._element.append(reference)


def add_run(p, text, font_key, red=False):
    for content, italic, _ in inline_spans(text):
        end = 0
        for match in REFERENCE_RE.finditer(content):
            if match.start() > end:
                add_plain_run(p, content[end:match.start()], font_key, red=red, italic=italic)
            add_reference(p, match.group(1) or match.group(2), font_key, red=red)
            end = match.end()
        if end < len(content) or not content:
            add_plain_run(p, content[end:], font_key, red=red, italic=italic)


def add_labeled(p, s, label_re):
    m = re.match(label_re, s)
    if m:
        add_run(p, m.group(1), "abstract_label")
        add_run(p, m.group(2), "abstract_body")
    else:
        add_run(p, s, "abstract_body")


first_content = True
for line in body_md.split("\n"):
    s = line.strip()
    if not s or s == "---":
        continue
    hashes = len(s) - len(s.lstrip("#"))
    if hashes and s[hashes:hashes + 1] == " ":
        s = s[hashes:].strip()
    else:
        hashes = 0
    if first_content and (hashes == 1 or (hashes == 0 and not s.startswith(("摘", "关键词")) and "作者姓名" not in s)):
        p = add_para(); p.alignment = 1
        add_run(p, s, "title")
        first_content = False
        continue
    first_content = False
    if s == "作者姓名":
        p = add_para(); p.alignment = 1
        add_run(p, s, "author")
    elif s.startswith("（单位名称"):
        p = add_para(); p.alignment = 1
        add_run(p, s, "affiliation")
    elif s.startswith(("摘　要", "摘要")):
        p = add_para(); p.alignment = 3
        add_labeled(p, s, r"^(摘\s*要[：:])(.*)$")
    elif s.startswith("关键词"):
        p = add_para(); p.alignment = 3
        add_labeled(p, s, r"^(关键词[：:])(.*)$")
    elif hashes == 2 or re.match(r"^[一二三四五六七八九十]+、", s):   # 一级标题
        p = add_para("一级标题")
        add_run(p, s, "h1")
    elif hashes == 3 or re.match(r"^（[一二三四五六七八九十]+）", s):  # 二级标题
        p = add_para("二级标题")
        add_run(p, s, "h2")
    elif hashes >= 4 or re.match(r"^\d+\.\s", s):                    # 三级标题
        p = add_para("三级标题")
        add_run(p, s, "h3")
    else:                                                               # 正文
        red_para = bool(RED_PARA_STARTS) and s.startswith(RED_PARA_STARTS)
        p = add_para("正文1")
        add_run(p, s, "body", red=red_para)

doc.save(TMP)

# ---- 注入真实脚注（支持红色标注） ----
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
def w(tag):
    return f"{{{W}}}{tag}"

def make_text_run(content, red=False, italic=False):
    r = etree.Element(w("r"))
    rPr = etree.SubElement(r, w("rPr"))
    fn_east, fn_west, fn_size = FONTS["footnote"]
    etree.SubElement(rPr, w("rFonts"), {w("ascii"): fn_west, w("hAnsi"): fn_west, w("eastAsia"): fn_east})
    half_points = str(int(round(fn_size * 2)))
    etree.SubElement(rPr, w("i"), {w("val"): "1" if italic else "0"})
    etree.SubElement(rPr, w("sz"), {w("val"): half_points})
    etree.SubElement(rPr, w("szCs"), {w("val"): half_points})
    if red:
        etree.SubElement(rPr, w("color"), {w("val"): "FF0000"})
    t = etree.SubElement(r, w("t"))
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = content
    return r


def make_text_runs(content, red=False, red_start=None):
    runs = []
    for text, italic, start in inline_spans(content):
        if red_start is None or red:
            runs.append(make_text_run(text, red=red, italic=italic))
        else:
            cut = max(0, min(len(text), red_start - start))
            if cut:
                runs.append(make_text_run(text[:cut], italic=italic))
            if cut < len(text):
                runs.append(make_text_run(text[cut:], red=True, italic=italic))
    return runs

fn_root = etree.Element(w("footnotes"), nsmap={"w": W})
sep = etree.SubElement(fn_root, w("footnote"), {w("type"): "separator", w("id"): "-1"})
sp = etree.SubElement(sep, w("p")); spp = etree.SubElement(sp, w("pPr"))
etree.SubElement(spp, w("spacing"), {w("after"): "0", w("line"): "240", w("lineRule"): "auto"})
etree.SubElement(etree.SubElement(sp, w("r")), w("separator"))
cont = etree.SubElement(fn_root, w("footnote"), {w("type"): "continuationSeparator", w("id"): "0"})
cp = etree.SubElement(cont, w("p")); cpp = etree.SubElement(cp, w("pPr"))
etree.SubElement(cpp, w("spacing"), {w("after"): "0", w("line"): "240", w("lineRule"): "auto"})
etree.SubElement(etree.SubElement(cp, w("r")), w("continuationSeparator"))

for fid in sorted(fn_map.keys()):
    fn = etree.SubElement(fn_root, w("footnote"), {w("id"): str(fid)})
    content = fn_map[fid]
    red = fn_sources[fid] in RED_FOOTNOTES
    red_start = (content.index(RED_SPLIT_MARK) if fn_sources[fid] == RED_FOOTNOTE_ID
                 and RED_SPLIT_MARK and RED_SPLIT_MARK in content else None)
    offset = 0
    for index, paragraph in enumerate(content.split("\n\n")):
        p = etree.SubElement(fn, w("p"))
        etree.SubElement(etree.SubElement(p, w("pPr")), w("pStyle"), {w("val"): FOOTNOTE_TEXT_STYLE.style_id})
        if index == 0:
            r1 = etree.SubElement(p, w("r"))
            etree.SubElement(etree.SubElement(r1, w("rPr")), w("rStyle"), {w("val"): FOOTNOTE_REF_STYLE.style_id})
            etree.SubElement(r1, w("footnoteRef"))
        local_start = red_start - offset if red_start is not None else None
        p.extend(make_text_runs(paragraph, red=red, red_start=local_start))
        offset += len(paragraph) + 2
footnotes_xml = etree.tostring(fn_root, xml_declaration=True, encoding="UTF-8", standalone=True)

zin = zipfile.ZipFile(TMP)
items = {}
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == "word/footnotes.xml":
        data = footnotes_xml
    items[item.filename] = (item, data)
zin.close()

# A minimal python-docx template may not already contain a footnotes part.
# Add the part plus its package relationship/content-type so the result remains
# a valid Word package instead of failing during post-generation verification.
if "word/footnotes.xml" not in items:
    items["word/footnotes.xml"] = (zipfile.ZipInfo("word/footnotes.xml"), footnotes_xml)

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

rels_name = "word/_rels/document.xml.rels"
if rels_name in items:
    rel_root = etree.fromstring(items[rels_name][1])
    if not any(node.get("Type") == REL_TYPE for node in rel_root):
        used_ids = {node.get("Id") for node in rel_root}
        rel_id = "rIdFootnotes"
        suffix = 2
        while rel_id in used_ids:
            rel_id = f"rIdFootnotes{suffix}"
            suffix += 1
        etree.SubElement(rel_root, f"{{{REL_NS}}}Relationship", {
            "Id": rel_id, "Type": REL_TYPE, "Target": "footnotes.xml"
        })
        items[rels_name] = (items[rels_name][0], etree.tostring(rel_root, xml_declaration=True, encoding="UTF-8"))

ct_name = "[Content_Types].xml"
if ct_name in items:
    ct_root = etree.fromstring(items[ct_name][1])
    footnote_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
    if not any(node.get("PartName") == "/word/footnotes.xml" for node in ct_root):
        etree.SubElement(ct_root, f"{{{CT_NS}}}Override", {
            "PartName": "/word/footnotes.xml", "ContentType": footnote_type
        })
        items[ct_name] = (items[ct_name][0], etree.tostring(ct_root, xml_declaration=True, encoding="UTF-8"))

with zipfile.ZipFile(DST, "x", zipfile.ZIP_DEFLATED) as zout:
    for fname, (info, data) in items.items():
        zout.writestr(info, data)

with zipfile.ZipFile(DST) as z:
    docxml = z.read("word/document.xml").decode("utf-8")
    refs = re.findall(r'<w:footnoteReference w:id="(\d+)"', docxml)
    red_body = docxml.count('w:val="FF0000"')
    fnxml = z.read("word/footnotes.xml").decode("utf-8")
    red_fn = fnxml.count('w:val="FF0000"')
expected_ids = {str(fid) for fid in fn_map}
actual_ids = set(re.findall(r'<w:footnote w:id="(\d+)"', fnxml))
if set(refs) != expected_ids:
    raise RuntimeError(f"脚注引用与Markdown不一致: refs={sorted(set(refs))}, expected={sorted(expected_ids)}")
if not expected_ids.issubset(actual_ids):
    raise RuntimeError(f"footnotes.xml 缺少脚注条目: missing={sorted(expected_ids - actual_ids)}")
print(f"生成: {DST} | 脚注引用: {len(refs)} 唯一: {len(set(refs))} | "
      f"正文红色run: {red_body} 脚注红色run: {red_fn} | 大小: {os.path.getsize(DST)}")
if REMOVE_TMP:
    try:
        os.unlink(TMP)
    except FileNotFoundError:
        pass
